from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from win32com.client import Dispatch


wb = Workbook()

sheet = wb['Sheet']

data = [('Month', 'Type', 'CRITICAL', 'HIGH', 'MEDIUM', 'LOW'), ('March 2022', 'Brute Force Login Attack Failed', '', '', 1, ''), ('', 'Inbound Communication With Blacklist IP Address', '', '', 5, ''), ('April 2022', 'Inbound Communication With Blacklist IP Address', '', '', 4, '')]
# insert data to excel
for item in data:
    sheet.append(item)

# declare chart ty Bar


chart1 = BarChart()
chart1.type = "col"
chart1.style = 1
chart1.title = "Bar Chart"
chart1.grouping = "stacked"
chart1.overlap = 100
chart1.y_axis.scaling.min = 0



# define chart bar
data = Reference(sheet, min_col=3, min_row=1, max_row=len(data), max_col=6)
cats = Reference(sheet, min_col=1, min_row=2, max_row=len(data), max_col=2)
chart1.add_data(data, titles_from_data=True)
chart1.set_categories(cats)
chart1.shape = 1

# colors lenend level
critical = chart1.series[0]
props = GraphicalProperties(solidFill="7030A0")
critical.graphicalProperties = props

high = chart1.series[1]
props = GraphicalProperties(solidFill="FF0000")
high.graphicalProperties = props

medium = chart1.series[2]
props = GraphicalProperties(solidFill="FFC000")
medium.graphicalProperties = props

low = chart1.series[3]
props = GraphicalProperties(solidFill="70AD47")
low.graphicalProperties = props

# set lenend position
chart1.legend.position = 'b'
chart1.dataLabels = DataLabelList()
chart1.dataLabels.showVal = True

# {'r', 'l', 'b', 'ctr', 'outEnd', 't', 'inBase', 'bestFit', 'inEnd'}
chart1.dataLabels.position = 'ctr'

# save chart
sheet.add_chart(chart1, "M2")

wb.save('sample.xlsx')

# save image
app = Dispatch("Excel.Application")
workbook_file_name = 'D:/INET-MS/Auto report/GitHub/WebReport/code_krit/report_SOC/sample.xlsx'

workbook = app.Workbooks.Open(Filename=workbook_file_name)

app.DisplayAlerts = False

i = 1
for sheet in workbook.Worksheets:
    for chartObject in sheet.ChartObjects():
        chartObject.Chart.Export('D:/INET-MS/Auto report/GitHub/WebReport/code_krit/report_SOC/' + str(i) + '.png')
        i += 1
workbook.Close(SaveChanges=False, Filename=workbook_file_name)

# os.system("sample.xlsx")