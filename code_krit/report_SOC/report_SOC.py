from ast import Global
from calendar import month
from logging import critical
from pickle import TRUE
from statistics import median
from turtle import circle
from executing import Source
import mysql.connector
from docxtpl import *
import calendar
from numpy import imag
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from win32com.client import Dispatch


mydb = mysql.connector.connect(
    host="10.11.101.32",
    user="admin01@INETMS",
    password="P@ssw0rd@INETMS",
    database="inetms_autoreport"
)


def count_event(data):
    count = 0
    for i in trend_DB:
        if i['priority'] == data:
            count += 1
    return count


def dictfetchall(cursor):
    # Return all rows from a cursor as a dict
    columns = [col[0] for col in cursor.description]
    return [
        dict(zip(columns, row))
        for row in cursor.fetchall()
    ]


def create_index(list, DB):
    temp_index = 1
    for i in DB:
        i['no'] = temp_index
        list.append(i)
        temp_index += 1
    return(list)


path = 'D:/INET-MS/Auto report/GitHub/WebReport/code_krit/report_SOC/'
doc = DocxTemplate(path+"templates/templateSOC.docx")
# =========================================================INPUT===========================================================
start_date = '2022-03-01'
end_date = '2022-04-01'
month = "มีนาคม"
year = "2565"
cus_name = "xspringam"

# ===========================================================Query=============================================================
mycursor = mydb.cursor()


sql = "select category,name_en,detail_th,priority from vulnerability"
mycursor.execute(sql)
vulnerability_DB = dictfetchall(mycursor)

sql = "select name,address,remark,type,name_th,service_provider from device join company_list ON device.company_id = company_list.id join device_type ON device_type.id = device.device_id where company_list.initials = \""+cus_name+"\""
mycursor.execute(sql)
device_DB = dictfetchall(mycursor)

sql = "select AVG(value) from capa_daily_log join company_list ON company_list.id = capa_daily_log.company_id where company_list.initials = \"" + \
    cus_name+"\" and capa_daily_log.date between \"" + \
    start_date+"\" and \""+end_date+"\""
mycursor.execute(sql)
capa_DB = dictfetchall(mycursor)
# category,type,action,des_address,port,protocol,src_address,des_geolocation,last_datetime,priority,des_geolocation,solution_th,file_path,src_user,src_geolocation,des_user,file_name,malware_name
sql = "select * from trend join vulnerability on vulnerability.name_en = trend.type join company_list ON company_list.id =trend.company_id where company_list.initials = \""+cus_name+"\"and trend.last_datetime between \"" + \
    start_date+"\" and \""+end_date+"\""
mycursor.execute(sql)
trend_DB = dictfetchall(mycursor)

temp = str(int(year)-543)
sql = "select type,priority,last_datetime from trend join vulnerability on vulnerability.name_en = trend.type join company_list ON company_list.id =trend.company_id where company_list.initials = \""+cus_name+"\"and trend.last_datetime between \"" + \
    temp+"-01-01\" and \""+temp+"-12-31\""
mycursor.execute(sql)
trend_graph = dictfetchall(mycursor)

# ================================================================Sum event=========================================================
crit = count_event('Critical (P1)')
high = count_event('High (P2)')
medium = count_event('Medium (P3)')
low = count_event('Low (P4)')
sumPriority = int(crit)+int(high)+int(medium)+int(low)
# =================================================================Device Table=====================================================
device = []
vulnerability = []
create_index(device, device_DB)
# =================================================================ภาคผนวก=========================================================
list_temp = []
for i in vulnerability_DB:
    i['sort'] = i['priority'][-2:-1]
    list_temp.append(i)
vulnerability_DB = sorted(list_temp, key=lambda d: d['name_en'])
vulnerability_DB = sorted(vulnerability_DB, key=lambda d: d['category'])
vulnerability_DB = sorted(vulnerability_DB, key=lambda d: d['sort'])
create_index(vulnerability, vulnerability_DB)

# ===================================================================detail==========================================================
head_tbl_vul_map = {'src_address': 'Source IP', 'des_address': 'Destination IP', 'src_geolocation': 'Source Geolocation', 'des_geolocation': 'Destination Geolocation', 'src_user': 'Source User', 'des_user': 'Destination User',
                    'port': 'Port', 'action': 'Event Subtype', 'last_datetime': 'Time', 'protocol': 'Protocol', 'url': 'URL', 'file_name': 'File Name', 'file_path': 'File Path', 'malware_name': 'Malware Name'}


head_tbl_vul = {'Inbound Communication With Blacklist IP Address': ['src_address', 'src_geolocation', 'des_address', 'port', 'action', 'last_datetime'],
                'Brute Force Login Attack Failed': [
                    'src_address', 'src_geolocation', 'src_user', 'des_address', 'last_datetime'],
                'Brute Force Login Attack Success': [
                    'src_address', 'src_geolocation', 'src_user', 'des_address', 'last_datetime'],
                'Outbound Communication with Blacklist IP Address Block': [
                    'src_address', 'des_address', 'des_geolocation', 'port', 'action', 'last_datetime'],
                'Outbound Communication with Blacklist IP Address Pass': [
                    'src_address', 'des_address', 'des_geolocation', 'port', 'action', 'last_datetime'],
                'SSH Access from Suspicious IP': [
                    'src_address', 'src_geolocation', 'des_address', 'action', 'last_datetime'],
                'RDP Access from Suspicious IP': [
                    'src_address', 'src_geolocation', 'des_address', 'action', 'last_datetime'],
                'Delegate Authentication Request': [
                    'src_address', 'src_user', 'des_address', 'des_user', 'action',  'last_datetime'],
                'SQL Injection (System Detect) Pass': [
                    'src_address', 'src_geolocation', 'des_address', 'action', 'last_datetime'],
                'Cross-Site Scripting (Custom Detected)': [
                    'src_address', 'src_geolocation', 'des_address', 'action', 'last_datetime'],
                'Virus or Spyware Detected': [
                    'src_address', 'des_address', 'des_geolocation', 'file_path', 'last_datetime'],
                'Virus or Spyware Detected But Failed to Clean': [
                    'src_address', 'des_address', 'des_geolocation', 'file_path', 'last_datetime'],
                'Port Sweep': [
                    'src_address', 'des_geolocation', 'port', 'last_datetime'],
                'Port Scan': ['src_address', 'des_geolocation', 'des_address', 'last_datetime']}

# =======================================================Vul Table==========================================================

dict_group_event = {}
list_category = []
for i in trend_DB:
    list_category.append(i['category'])

    if i['type'] in dict_group_event:
        dict_group_event[i['type']] += 1
    else:
        dict_group_event[i['type']] = 1
list_category = list(dict.fromkeys(list_category))

new_dict_group_event = dict((k.lower(), v)
                            for k, v in dict_group_event .items())

list_group_category = []
for j in list_category:
    dict_group_category = {}
    list_event = []
    count = 0
    for i in vulnerability_DB:
        if j == i['category']:
            if i['name_en'].lower() in new_dict_group_event:
                event = {}
                event['type'] = i['name_en']
                event['priority'] = i['priority']
                event['count'] = new_dict_group_event[i['name_en'].lower()]
                count += new_dict_group_event[i['name_en'].lower()]
                list_event.append(event)
    dict_group_category['category'] = j
    dict_group_category['count'] = count
    dict_group_category['event'] = list_event
    list_group_category.append(dict_group_category)

count = 0
for i in list_group_category:
    count += i['count']

# ==========================================================Event Detail==========================================================

all_event_list = []
head_tbl_vul_new = []

for x in dict_group_event:
    one_event_dict = {}
    one_event_list = []

    for i in trend_DB:
        event = []
        event_dict = {}
        if x == i['type']:
            for j in head_tbl_vul[x]:
                event.append(i[j])
            event_dict['event'] = event
            one_event_list.append(event_dict)
            one_event_dict['cols'] = one_event_list
            one_event_dict['detail'] = i['solution_th']
            one_event_dict['vul'] = i['intro_th']

    head_tbl_vul_new = []
    for j in head_tbl_vul[x]:
        head_tbl_vul_new.append(head_tbl_vul_map[j])
    one_event_dict['col_label'] = head_tbl_vul_new
    all_event_list.append(one_event_dict)

# ==================================================Gen Graph======================================================
wb = Workbook()
sheet = wb['Sheet']
month_graph = []
data_for_graph = [('Month', 'Type', 'CRITICAL', 'HIGH', 'MEDIUM', 'LOW')]
dict_graph = {}
dict_list_graph = {}
for i in trend_graph:
    temp = (str(i['last_datetime'])[0:7].split('-'))
    month_graph = (calendar.month_name[int(temp[1])]) + ' ' + temp[0]
    if month_graph not in dict_graph:
        dict_graph[month_graph] = {}
    if i['type'] not in dict_graph[month_graph]:
        dict_graph[month_graph][i['type']] = {
            'CRITICAL': 0, 'HIGH': 0, 'MEDIUM': 0, 'LOW': 0}

    if i['priority'] == 'critical (P1)':
        dict_graph[month_graph][i['type']]['CRITICAL'] += 1
    elif i['priority'] == 'High (P2)':
        dict_graph[month_graph][i['type']]['HIGH'] += 1
    elif i['priority'] == 'Medium (P3)':
        dict_graph[month_graph][i['type']]['MEDIUM'] += 1
    elif i['priority'] == 'Low (P4)':
        dict_graph[month_graph][i['type']]['LOW'] += 1
for i, j in dict_graph.items():
    check = TRUE
    for x, y in j.items():
        cri_graph = '' if y['CRITICAL'] == 0 else y['CRITICAL']
        high_graph = '' if y['HIGH'] == 0 else y['HIGH']
        med_graph = '' if y['MEDIUM'] == 0 else y['MEDIUM']
        low_graph = '' if y['LOW'] == 0 else y['LOW']

        if check:
            temp = (i, x, cri_graph, high_graph, med_graph, low_graph)
            check = False
        else:
            temp = ('', x, cri_graph, high_graph, med_graph, low_graph)
        data_for_graph.append(temp)

for item in data_for_graph:
    sheet.append(item)

# declare chart ty Bar
chart1 = BarChart()
chart1.type = "col"
chart1.style = 1
chart1.title = "Incident Summary "+str(int(year)-543)
chart1.grouping = "stacked"
chart1.overlap = 100
chart1.y_axis.scaling.min = 0

# define chart bar
data_for_graph = Reference(sheet, min_col=3, min_row=1,
                           max_row=len(data_for_graph), max_col=6)
cats = Reference(sheet, min_col=1, min_row=2,
                 max_row=len(data_for_graph), max_col=2)
chart1.add_data(data_for_graph, titles_from_data=True)
chart1.set_categories(cats)
chart1.shape = 1

# # colors lenend level
critical_graph = chart1.series[0]
props = GraphicalProperties(solidFill="7030A0")
critical_graph.graphicalProperties = props

high_graph = chart1.series[1]
props = GraphicalProperties(solidFill="FF0000")
high_graph.graphicalProperties = props

medium_graph = chart1.series[2]
props = GraphicalProperties(solidFill="FFC000")
medium_graph.graphicalProperties = props

low_graph = chart1.series[3]
props = GraphicalProperties(solidFill="70AD47")
low_graph.graphicalProperties = props

# # set lenend position
chart1.legend.position = 'b'
chart1.dataLabels = DataLabelList()
chart1.dataLabels.showVal = True

# # {'r', 'l', 'b', 'ctr', 'outEnd', 't', 'inBase', 'bestFit', 'inEnd'}
chart1.dataLabels.position = 'ctr'

# # save chart
sheet.add_chart(chart1, "M2")
wb.save('sample.xlsx')

# save image
app = Dispatch("Excel.Application")
workbook_file_name = path+'sample.xlsx'
workbook = app.Workbooks.Open(Filename=workbook_file_name)
app.DisplayAlerts = False


for sheet in workbook.Worksheets:
    for chartObject in sheet.ChartObjects():
        chartObject.Chart.Export(path + 'graph.png')

workbook.Close(SaveChanges=False, Filename=workbook_file_name)

# ========================================================Content===============================================
contents = {}
img = "inet.png" if device_DB[0]['service_provider'] == "INET" else "inetms.png"
image = InlineImage(doc, path+img)
contents['img_provider'] = image
contents['month'] = month
contents['year'] = year
contents['Service_provider'] = "บริษัท อินเทอร์เน็ตประเทศไทย จำกัด (มหาชน)" if device_DB[
    0]['service_provider'] == "INET" else "บริษัท ไอเน็ต แมเนจด์ เซอร์วิสเซส จำกัด"
contents['name_capany'] = device_DB[0]['name_th']
contents['capa'] = '{:.4f}'.format(float(capa_DB[0]['AVG(value)']))
contents['vulnerability'] = vulnerability
contents['device'] = device
contents['crit'] = crit
contents['medium'] = medium
contents['low'] = low
contents['sum'] = sumPriority
contents['high'] = high
contents['all_event'] = all_event_list
contents['category'] = list_group_category
contents['category_count'] = count
graph = InlineImage(doc, path+'graph.png')
contents['graph'] = graph
doc.render(contents)
doc.save(path+"output\\รายงานสรุปผลการปฏิบัติงานของศูนย์SOC_" +
         cus_name+'_'+month+'_'+year+".docx")
# os.system(path+"SOC.docx")
